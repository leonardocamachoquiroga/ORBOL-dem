from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'OLBOL propuesta y guion de demo.docx'
SHOTS = ROOT / 'docs' / 'capturas'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.header_distance, sec.footer_distance = Inches(.28), Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.13
doc.styles['Normal'].paragraph_format.space_after = Pt(9)
doc.styles['Title'].font.size = Pt(30)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(13)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(22)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(7)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(12)
header = sec.header.paragraphs[0]
header.text = 'OLBOL   /   PROPUESTA COMERCIAL Y GUION DE PRESENTACIÓN'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0,0,0)
footer = sec.footer.paragraphs[0]
footer.text = 'ZEKIRI para OLBOL   ·   2 de octubre de 2026'
footer.runs[0].font.size = Pt(8)
footer.add_run(' '*10 + 'Página ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title = 'OLBOL propuesta y guion de demo'
doc.core_properties.subject = 'WhatsApp oficial CRM básico costos y presentación comercial'
doc.core_properties.author = 'ZEKIRI'

def p(text='',bold=False):
    para=doc.add_paragraph()
    run=para.add_run(text);run.bold=bold
    return para
def h(text): doc.add_heading(text,2)
def page(title): doc.add_page_break();doc.add_heading(title,1)
def shot(name, caption, width=6.9):
    path=SHOTS/name
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(5)
    para.add_run().add_picture(str(path),width=Inches(width))
    inline=para.runs[0]._r.xpath('.//wp:docPr')[0]
    inline.set('descr',caption)
    doc.add_paragraph(caption,'Caption')
def bullet(text):
    para=doc.add_paragraph(text,'List Bullet')
    para.paragraph_format.space_after=Pt(6)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers))
    t.autofit=False
    for i,heading in enumerate(headers):t.rows[0].cells[i].text=heading
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=value
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'27303B' if ri==0 else ('F0F3F6' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D7DDE3');borders.append(border)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),'100');node.set(qn('w:type'),'dxa');margins.append(node)
            tcPr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(10)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        trPr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p()
    return t
def link(title,url):
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(7)
    hyperlink=OxmlElement('w:hyperlink')
    rid=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink.set(qn('r:id'),rid)
    r=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'295A7A');props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');props.append(size)
    r.append(props);text=OxmlElement('w:t');text.text=title;r.append(text);hyperlink.append(r);para._p.append(hyperlink)

doc.add_paragraph('OLBOL propuesta y guion de demo','Title')
doc.add_paragraph('Preparado por ZEKIRI para presentar al dueño de OLBOL. Documento de apoyo para el equipo que conduce la reunión. Investigación consultada el 2 de octubre de 2026.','Subtitle')
p('Proponemos un piloto que conecte la experiencia web, WhatsApp oficial y un CRM básico. El objetivo es que cada consulta llegue a un vendedor con contexto y un siguiente paso, y que OLBOL pueda medir qué campañas generan oportunidades.',True)
shot('inicio-escena.png','Captura de la demo revisada. Los vehículos, precios, sucursales y testimonios son conceptuales.',width=5.7)
p('La decisión a buscar en la reunión es aprobar el alcance del piloto y designar un responsable de OLBOL. Los honorarios incluidos son una propuesta sugerida para nuestro equipo; deben confirmarse antes de emitir la oferta formal.')
table(['Punto de partida','Propuesta'],[
    ('Implementación sugerida','USD 2.400 en tres hitos'),
    ('Con anuncios de USD 450','USD 885–945 al mes en total estimados'),
], [3.3,3.6])

page('Preparación y apertura de la reunión')
h('Antes de presentar')
bullet('Abrir la demo local y precargar Inicio, Vehículos, Terra S5, Asesor, CRM demo y Propuesta. Usar zoom de navegador al 100% y una conexión estable.')
bullet('En Asesor pulsar Reiniciar demo. En CRM pulsar Restablecer ejemplos. Tener las capturas del documento disponibles si falla el equipo o la red.')
bullet('Usar datos ficticios en formularios. El CRM de esta demo guarda ejemplos únicamente en el navegador; no representa una integración ya contratada.')
h('Preguntas de diagnóstico durante la apertura')
p('Antes de enseñar pantallas, preguntar al dueño cuántas consultas recibe por semana, quién responde hoy y qué ocurre cuando un interesado no compra en el primer contacto. Confirmar si su número ha recibido restricciones y cómo consigue permiso para el seguimiento. Escuchar las respuestas y conectar la demostración con ese problema concreto.')
h('Speech de apertura de 45 segundos')
p('“Queremos mostrarle cómo podría vender OLBOL cuando el catálogo, las consultas y el equipo trabajan conectados. Un cliente llega desde un anuncio o la web, recibe orientación y avanza a una cotización o una visita. Su vendedor recibe esa conversación con contexto y sabe qué hacer después. Hoy verá el recorrido en una demo; después le explicaremos la integración oficial y la inversión para un piloto.”')
h('Recorrido propuesto de 10 a 12 minutos')
table(['Momento','Pantalla','Qué demostrar'],[
    ('0–2 min','Inicio y catálogo','Confianza y exploración sencilla'),
    ('2–4 min','Ficha y asesor','Recomendación basada en necesidades'),
    ('4–6 min','Comparación y cotización','Avance hacia una decisión'),
    ('6–8 min','WhatsApp y CRM','Continuidad y responsable comercial'),
    ('8–12 min','Propuesta','Costos alcance y siguiente paso'),
], [1,1.7,4.2])

page('Inicio y catálogo')
shot('catalogo-escena.png','Catálogo de demostración con filtros, imágenes y precio visible.')
h('Acción en pantalla')
p('Abrir Inicio, entrar en Vehículos y filtrar SUV. Abrir Terra S5. Si se desea, probar una búsqueda sin resultados y limpiar los filtros para mostrar la recuperación.')
h('Speech')
p('“La primera impresión tiene que estar a la altura de OLBOL. El vehículo es protagonista, el precio está visible y el cliente puede explorar sin perderse. Esta experiencia también funciona en móvil. La información de hoy es conceptual; al arrancar reemplazamos catálogo, fotografías, condiciones y precios por los que ustedes aprueben.”')
h('Valor para el dueño')
p('El catálogo funciona como un punto de entrada comercial. Permite pasar de una curiosidad general a un modelo concreto y continuar con el asesor sin obligar al comprador a completar un formulario largo.')

page('Ficha del vehículo y asesor')
shot('vehiculo-escena.png','Ficha Terra S5 con galería ampliable, datos del catálogo y acciones comerciales.')
h('Acción en pantalla')
p('Ampliar una imagen y cerrar con Escape. Mostrar las especificaciones y entrar a Cotizar con el asesor. En el asesor escribir: “Somos 5, viajamos los fines de semana y tengo un presupuesto de USD 30 mil”.')
h('Speech')
p('“Aquí el comprador puede revisar el vehículo y preguntar si encaja con su vida. El asesor recoge necesidades como presupuesto, pasajeros y uso. Con ese contexto le propone una opción. La interfaz usa datos del catálogo y el simulador calcula las cuotas; para producción gobernamos el catálogo y las condiciones con su equipo.”')
p('El asesor local mantiene la demostración disponible si falla la llamada al servicio externo. La activación de IA externa y su calidad se validan durante el piloto; no debe decidir descuentos, disponibilidad real ni aprobación bancaria sin autorización.')

page('Recomendación y comparación')
shot('recomendacion-escena.png','El asesor recomienda Terra S5 para el escenario familiar de la presentación.')
h('Acción en pantalla')
p('Mostrar la recomendación, enviar “Quiero comparar” y revisar las diferencias. Después escribir “Quiero cotizar”. La pantalla cambia del comparador al simulador.')
h('Speech')
p('“El cliente no necesita saber de antemano qué modelo comprar. Nos cuenta lo que necesita y ve una recomendación explicada. Si tiene dudas, puede comparar. Queremos que avance con información y que el vendedor reciba esas prioridades, en lugar de empezar otra conversación desde cero.”')
h('Frase para conectar con la operación')
p('“¿Cuánto tiempo le toma hoy a su equipo volver a preguntar presupuesto, uso y modelo de interés? Aquí ese contexto acompaña a la oportunidad.”')

page('Cotización y continuidad por WhatsApp')
shot('cotizacion-escena.png','Simulación de cuota y mensaje preparado con el contexto de la conversación.',width=4.7)
h('Acción en pantalla')
p('Mover la inicial, alternar 36 y 48 meses, validar la cotización y pulsar Preparar continuidad. Mostrar el resumen y, si se desea, copiarlo. Aclarar que la demo no envía mensajes ni crea reservas reales.')
h('Speech')
p('“La conversación llega a un siguiente paso concreto. El comprador ajusta una simulación y puede seguir por WhatsApp con el modelo y sus necesidades. En producción conectamos el número oficial, las plantillas y el CRM. La tasa de 10,5% anual es un ejemplo del simulador; la oferta financiera real debe venir de las condiciones aprobadas por OLBOL y su entidad financiera.”')
p('Después enviar “Me interesa” y abrir Ver CRM. Ese interés aparece como Visitante web demo, con una tarea pendiente de asignación y permiso de contacto.')

page('CRM básico para el equipo comercial')
shot('crm-escena.png','Consola local de demostración. Las cifras y personas no son resultados reales de OLBOL.')
h('Acción en pantalla')
p('Abrir la oportunidad Visitante web demo, asignarla a Sofía, moverla a Prueba y guardar una nota. Mostrar responsable, próxima acción, historial y estado del permiso de seguimiento.')
h('Speech')
p('“Esta es la parte que evita que una buena consulta quede olvidada. Cada oportunidad tiene un responsable y un próximo paso. El vendedor ve qué buscaba el comprador y qué se conversó. Usted puede revisar qué llegó de anuncios, qué llegó de la web y cuánto avanzó. Para el piloto proponemos usar Kommo como CRM; esta consola ilustra la operación, no replica toda su plataforma.”')
h('Alcance básico del piloto')
p('Contactos y oportunidades, origen de campaña, etapas, responsables, notas, tareas, historial y permisos. El informe inicial medirá primera respuesta, oportunidades calificadas, pruebas de manejo y cierres por origen. Se define una línea de base antes de prometer mejoras.')

page('Solución de WhatsApp para ventas y anuncios')
p('Recomendamos WhatsApp Business Platform con la integración oficial de Kommo y anuncios Click to WhatsApp. La API oficial es la vía de integración; su uso no elimina la posibilidad de restricciones por parte de Meta. [1] [2]',True)
h('Cómo reducimos el riesgo de restricciones')
p('La operación registra el permiso de contacto y distingue seguimiento de una solicitud de promociones. Se usan plantillas aprobadas cuando corresponde y se respeta la ventana de atención de 24 horas. Las solicitudes de baja detienen el seguimiento. La automatización ofrece una vía clara para hablar con una persona. [1]')
p('Proponemos asignar un responsable de revisar calidad y restricciones en WhatsApp Manager. Si aparecen señales de deterioro, se pausan envíos salientes y se revisan público, frecuencia, permisos y contenido. No usamos herramientas de automatización de WhatsApp Web, bases compradas ni rotación de números para eludir sanciones.')
h('Qué hacer con el número actual')
p('Revisar titularidad, acceso, estado y antecedentes del número antes de integrar. Kommo documenta coexistencia de WhatsApp Business App y la plataforma para números elegibles. Si esa vía no está disponible, se acuerda una migración y una ventana de cambio con OLBOL antes de modificar el canal. [3]')
h('Cómo usamos los anuncios')
p('El anuncio invita al comprador a escribir. Registramos campaña, anuncio y modelo de interés. Responder a una entrada elegible de Click to WhatsApp dentro de 24 horas puede abrir una ventana de 72 horas sin cargos de Meta según las condiciones aplicables. La pauta y los posibles cargos de proveedor siguen separados. Esa ventana comercial no es un permiso para enviar mensajes libres sin respetar la ventana de atención. [2] [4]')
h('Speech para responder sobre bloqueos')
p('“No le vamos a prometer una cuenta que Meta nunca pueda restringir. Le proponemos trabajar con la integración oficial, clientes que autoricen el contacto y una operación que revise la calidad antes de aumentar envíos. También dejamos un canal humano y una forma sencilla de dejar de recibir mensajes.”')

page('Alternativas y recomendación del CRM')
p('Nuestra recomendación inicial es Kommo Advanced: combina embudo comercial y automatizaciones a un costo por usuario comprensible. Evaluar la tarifa regional y confirmar en una prueba que su integración cubre el número y las funciones que necesita OLBOL. [5]')
table(['Opción','Costo publicado de referencia','Encaje y condición'],[
    ('Kommo Base','USD 25 por usuario al mes\n3 usuarios USD 75','CRM e inbox básicos. Para tareas manuales. Contratación mínima de 6 meses.'),
    ('Kommo Advanced\nRecomendado','USD 35 por usuario al mes\n3 usuarios USD 105','Añade automatizaciones. Licencia inicial de 3 usuarios por 6 meses: USD 630.'),
    ('respond.io Growth','USD 199 al mes mensual\nUSD 159 equivalente anual','10 usuarios y nivel inicial de 1.000 contactos activos. Incluye flujos; el consumo escala. Más orientado a operación de conversaciones.'),
    ('Chatwoot Business','USD 39 por agente al mes\n3 agentes USD 117','Inbox y reglas de automatización. Si se requiere un embudo comercial, hay que configurarlo o integrarlo aparte.'),
], [1.45,2.35,3.1])
p('Las opciones no se suman entre sí. Elegimos un proveedor principal. Los cargos de Meta, anuncios y consumo adicional quedan fuera de estas licencias. Respond.io ofrece otra estructura de usuarios y contactos; revisar límites y créditos antes de comparar solo el precio. [5] [6] [7]')
h('API directa como alternativa')
p('Conectar Cloud API directamente daría más control y exigiría construir operación, seguimiento y soporte. Twilio publica un recargo de USD 0,005 por mensaje entrante o saliente, además de los cargos de Meta; 10.000 mensajes representarían USD 50 de recargo Twilio. No se recomienda sumar Twilio a la integración de Kommo por defecto. [8]')
h('Decisión propuesta')
p('Empezar con Kommo Advanced y tres usuarios. Pasar a una solución propia solo si el volumen o los requisitos del piloto justifican el costo adicional de desarrollo y operación.')

page('Inversión y condiciones de la propuesta')
p('Escenario orientativo de un número, tres usuarios y USD 450 mensuales de anuncios. Los honorarios son sugeridos; los importes de hosting, IA, dominio y mensajería son provisiones de planificación, no tarifas ni límites garantizados.')
table(['Concepto','Importe en USD','Cómo se trata'],[
    ('Implementación','2.400 una vez','Honorario sugerido por alcance'),
    ('Kommo Advanced','105 al mes equivalente','630 por 6 meses al contratar [5] [9]'),
    ('Hosting web','25 al mes','Provisión estimada'),
    ('Consumo de IA','25 al mes','Provisión estimada'),
    ('WhatsApp y Meta','30–90 al mes','Provisión variable por uso'),
    ('Tecnología mensual','185–245','Suma CRM hosting IA y WhatsApp'),
    ('Mantenimiento nuestro','250 al mes','Hasta 3 h de ajustes menores'),
    ('Publicidad en Meta','450 al mes','Inversión del cliente en anuncios'),
    ('Total mensual','885–945','Tecnología mantenimiento y anuncios'),
    ('Dominio','25 al año','Provisión separada'),
], [2.5,1.8,2.6])
p('Provisión de arranque: USD 3.835–3.895, si se reserva toda la implementación, seis meses de CRM, dominio y el primer mes de operación. No se duplica la licencia CRM mensual en esta suma. Presupuesto total a seis meses: USD 7.735–8.095 incluyendo implementación y dominio.')
p('Pago de implementación sugerido: USD 1.200 al inicio, USD 720 al completar las integraciones y USD 480 tras la aceptación. Licencias y publicidad se pagan desde cuentas de OLBOL. Importes sin impuestos, comisiones de pago, excesos de consumo ni nuevas funciones.')
p('El mantenimiento incluye monitoreo básico, revisión mensual y hasta tres horas de ajustes menores. No incluye atención 24/7, producción continua de piezas ni gestión diaria de anuncios. Esta última puede cotizarse como servicio adicional; estimación interna sugerida de USD 150 al mes, fuera de los totales.')

page('Tarifas de mensajes y validación antes de contratar')
p('El costo de WhatsApp depende del país del destinatario, categoría del mensaje y exenciones aplicables. No existe una tarifa única por contacto. La provisión de USD 30–90 sirve para planificar y debe reemplazarse por una estimación con volumen y tarifario confirmados. [2]')
h('Cambio reciente que afecta el presupuesto')
p('Kommo publicó el 30 de septiembre de 2026 que desde el 1 de octubre se cobran mensajes de servicio después de los primeros 1.000 por número al mes, y mensajes de utilidad dentro de la ventana de atención salvo una exención aplicable. Reporta que los mensajes elegibles durante la ventana gratuita de 72 horas siguen exentos. [4]')
p('La página general de Meta consultada todavía describe respuestas de atención y utilidad sin cargo. La documentación técnica de tarifas de Meta no pudo consultarse por un error de acceso 429. Por esa discrepancia no presentamos un precio unitario de Bolivia como definitivamente validado ni prometemos respuestas ilimitadas gratuitas. [2]')
h('Validación que cierra el presupuesto de consumo')
bullet('Confirmar tarifa vigente para destinatarios de Bolivia y, si corresponde, otros países; categoría y fecha de vigencia en WhatsApp Manager o con el proveedor.')
bullet('Estimar mensajes de servicio, utilidad y marketing entregados fuera de exenciones. Separar entradas de anuncios elegibles y conversaciones desde la web.')
bullet('Multiplicar mensajes facturables por su tarifa, aplicar tramos y añadir cualquier recargo del proveedor. Reservar margen por tráfico mayor al previsto.')
bullet('Definir un aviso de gasto y una revisión de presupuesto al alcanzar el 80% de la provisión. No bloquear la atención del cliente sin acordar un procedimiento.')
h('Ejemplo comercial para explicar el valor')
p('Si una venta adicional aportara USD 1.000 de margen de contribución, el costo mensual orientativo de USD 885–945 equivaldría a cerca de una venta adicional. Es un ejemplo aritmético, no un pronóstico. OLBOL debe aportar su margen real, tasa de conversión y costos de venta para evaluar retorno.')

page('Alcance del piloto y aceptación')
h('Entrega propuesta en cuatro a seis semanas')
table(['Etapa','Entrega verificable'],[
    ('Semana 1','Catálogo aprobado permisos activos del negocio y embudo definidos'),
    ('Semanas 2 y 3','Web asesor WhatsApp oficial y CRM conectados'),
    ('Semana 4','Campaña piloto inicial capacitación y medición de consultas'),
    ('Semanas 5 y 6','Correcciones prueba completa y aceptación del piloto'),
],[1.8,5.1])
p('El calendario comienza con la entrega de accesos y materiales. La verificación de empresa, elegibilidad del número y aprobación de plantillas dependen de Meta y pueden desplazar fechas. El alcance no incluye pagos, reserva legal, crédito bancario, migraciones extensas ni un CRM propio de producción.')
h('Criterios de aceptación')
bullet('Un cliente de prueba llega desde la web o campaña y genera una oportunidad con origen, vehículo y contexto correctos.')
bullet('La bandeja oficial recibe el mensaje y permite intervención humana. Un seguimiento autorizado utiliza una plantilla aprobada cuando corresponde.')
bullet('El vendedor puede asignar responsable, etapa y próxima acción; la solicitud de baja bloquea seguimiento comercial.')
bullet('El catálogo y los precios de producción coinciden con los materiales aprobados por OLBOL. Las simulaciones financieras se identifican como orientativas.')
bullet('El equipo recibe capacitación y se entrega un reporte de respuesta, oportunidades y citas del piloto.')
h('Qué necesitamos de OLBOL')
p('Responsable comercial, acceso administrativo seguro al portafolio de Meta, control del número, catálogo y precios oficiales, fotografías autorizadas, información de sucursales, horarios, políticas comerciales, privacidad y ejemplos de mensajes. Acordar los permisos mediante roles; las credenciales no se incluyen en documentos ni Git.')

page('Cierre comercial y respuestas a objeciones')
h('Speech de cierre de un minuto')
p('“Ya vimos cómo un interesado puede descubrir un modelo, resolver dudas y avanzar a una cotización. Ahora su equipo puede recibir ese interés con contexto y una tarea concreta. Nuestra propuesta es empezar con un piloto: un número oficial, tres usuarios del CRM y una campaña inicial. Planteamos USD 2.400 de implementación y USD 250 mensuales de mantenimiento, más los costos de tecnología y anuncios que acabamos de separar. Si el alcance le encaja, hoy podemos acordar un responsable, los materiales necesarios y la fecha de inicio. Le entregamos la oferta formal con hitos y criterios de aceptación.”')
h('El cierre que buscamos')
p('“¿Podemos acordar hoy el alcance del piloto y quién de OLBOL lo liderará con nosotros?” Después de la respuesta, dejar por escrito alcance, responsables, precio formal y próximos hitos. No dar por aceptados los honorarios internos solo por haber mostrado la demo.')
h('Si pregunta por el costo')
p('“Podemos reducir el número de usuarios o empezar con menos pauta. Lo que queremos conservar es un recorrido medible y un responsable por oportunidad. Primero validamos el costo de perder o demorar consultas en su operación.”')
h('Si pregunta por la IA o por reemplazar a sus vendedores')
p('“El asesor cubre orientación inicial y organiza contexto. Su equipo conserva el cierre, la negociación, la disponibilidad real y los casos que requieren criterio comercial. En el piloto revisamos las conversaciones y corregimos el comportamiento antes de escalar.”')
h('Si pide garantías de ventas o de no bloqueo')
p('“Podemos comprometernos con la entrega, el proceso y los criterios de aceptación. Las ventas dependen también de oferta, inventario, atención y mercado. Meta conserva sus políticas de restricción. Medimos el piloto y tomamos decisiones con sus resultados.”')
h('Si no desea decidir en la reunión')
p('Acordar una revisión de la oferta con fecha y participantes. Pedir los datos pendientes que modifican el presupuesto: número de vendedores, volumen de consultas, estado del WhatsApp y presupuesto de anuncios. Mantener el siguiente paso concreto sin presionar con urgencia inventada.')

page('Fuentes de la investigación')
p('Fuentes primarias de las plataformas y proveedores, consultadas el 2 de octubre de 2026. Las condiciones comerciales pueden variar por región, cuenta y fecha de contratación.')
sources=[
    ('[1] WhatsApp Business Messaging Policy','https://business.whatsapp.com/policy'),
    ('[2] Meta WhatsApp Business Platform Pricing','https://business.whatsapp.com/products/platform-pricing'),
    ('[3] Kommo WhatsApp Business y coexistencia','https://support.kommo.com/docs/es/whatsapp-business-overview'),
    ('[4] Kommo cambio de precios de WhatsApp octubre de 2026','https://www.kommo.com/blog/whatsapp-pricing-update/'),
    ('[5] Kommo tarifas oficiales','https://www.kommo.com/buy/tariff/'),
    ('[6] respond.io planes y precios','https://respond.io/pricing'),
    ('[7] Chatwoot planes y precios','https://www.chatwoot.com/pricing'),
    ('[8] Twilio WhatsApp Messaging Pricing','https://www.twilio.com/en-us/whatsapp/pricing'),
    ('[9] Kommo contratación y facturación','https://support.kommo.com/docs/billing-overview'),
]
for title,url in sources:link(title,url)
h('Notas sobre evidencia y estimaciones')
p('Los precios de licencias provienen de páginas oficiales. El precio de implementación, mantenimiento y gestión adicional corresponde a una propuesta interna sugerida, no a una tarifa de proveedor. Las provisiones de infraestructura y consumo son estimaciones de planificación. La discusión sobre la actualización de WhatsApp distingue el aviso de Kommo de la página general de Meta y deja pendiente confirmar el tarifario técnico aplicable.')
p('Todas las capturas se tomaron de la demo local revisada. Las métricas y oportunidades del CRM son ficticias. El guion comercial propone cómo conducir la reunión y no representa resultados, promesas ni acuerdos ya aceptados por OLBOL.')

doc.save(OUT)
print(OUT)
