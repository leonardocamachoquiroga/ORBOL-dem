from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'OLBOL WhatsApp IA y Yaku.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption', 'List Bullet']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for node in style.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Title'].font.size = Pt(29)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Subtitle'].font.size = Pt(11)
doc.styles['Heading 1'].font.size = Pt(23)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(10)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('OLBOL  |  2 de octubre de 2026  |  ').font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'OLBOL WhatsApp IA y Yaku'
doc.core_properties.subject = 'Opciones requisitos costos y guion para decidir la integración comercial'
doc.core_properties.author = 'ZEKIRI'


def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para


def h(text):
    doc.add_heading(text, 2)


def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def bullet(text):
    para = doc.add_paragraph(text, 'List Bullet')
    para.paragraph_format.space_after = Pt(5)


def table(headers, rows, widths, centered=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        cells = t.add_row().cells
        for i, text in enumerate(values):
            cells[i].text = text
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '30363C' if ri == 0 else ('F3F4F5' if ri % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    node.set(qn('w:' + key), value)
                borders.append(node)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '105'), ('bottom', '105'), ('left', '115'), ('right', '115')]:
                node = OxmlElement('w:' + side)
                node.set(qn('w:w'), value)
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            props.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                if ci in centered:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(10)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1
    para.add_run().font.size = Pt(3)


def sources(items):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(8)
    for i, (title, url) in enumerate(items):
        if i:
            para.add_run('  |  ').font.size = Pt(9)
        node = OxmlElement('w:hyperlink')
        node.set(qn('r:id'), para.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
        color = OxmlElement('w:color'); color.set(qn('w:val'), '295A7A'); props.append(color)
        run.append(props)
        text = OxmlElement('w:t'); text.text = title; run.append(text)
        node.append(run); para._p.append(node)


def shot(name, caption, width=5.8):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run()
    run.add_picture(str(ROOT / 'docs' / 'capturas' / name), width=Inches(width))
    run._r.xpath('.//wp:docPr')[0].set('descr', caption)
    doc.add_paragraph(caption, 'Caption')


doc.add_paragraph('OLBOL WhatsApp IA y Yaku', 'Title')
doc.add_paragraph('Comparación para definir la solución comercial y presentar el piloto al dueño de OLBOL. Preparado por ZEKIRI. Investigación al 2 de octubre de 2026.', 'Subtitle')
p('Proponemos Yaku como CRM y un asesor de IA conectado al catálogo mediante WhatsApp Business Platform. La primera ruta a evaluar es Meta Cloud API directa: evita una cuota adicional de intermediario, pero exige que mantengamos la integración y que Yaku pueda recibir conversaciones y oportunidades.', True)
p('Si OLBOL prefiere un CRM propio, cotizaremos uno básico dentro del panel administrativo del DMS. La elección del CRM es independiente del proveedor de WhatsApp: ambas opciones pueden utilizar Meta directa, Twilio o 360dialog.')
shot('crm-escena.png', 'Consola de la demo con datos ficticios. Ilustra el seguimiento propuesto; no es una captura de Yaku ni una conexión real a WhatsApp.', 5.7)
h('Decisión recomendada')
p('Validar primero las capacidades de Yaku y el número actual. Elegir 360dialog si conservar la app mediante coexistencia elegible y el acompañamiento de alta justifican su cuota. Evaluar Twilio si ya usamos su infraestructura o el volumen inicial es pequeño. El piloto debe demostrar atención, seguimiento y control humano antes de ampliar campañas.')

page('Opciones de conexión oficial')
p('Las tarifas de esta tabla corresponden al acceso o transporte de un número. El consumo de Meta, nuestra IA, el CRM, infraestructura y honorarios se calculan aparte.')
table(['Opción', 'Costo del proveedor', 'Cuándo la proponemos'], [
    ('Meta Cloud API directa', 'Sin cuota adicional de acceso\nMás consumo de Meta', 'Ruta preferida si podemos operar el backend y conectar Yaku o el DMS.'),
    ('Twilio Programmable Messaging', 'USD 0,005 por mensaje entrante o saliente\nMás consumo de Meta', 'Integración con infraestructura Twilio existente o piloto con poco volumen.'),
    ('360dialog Regular', 'USD 59 por número al mes\nMás consumo de Meta', 'Alta acompañada y coexistencia de app y API para un número elegible.'),
    ('App WhatsApp Business sola', 'App gratuita\nPublicidad y extras aparte', 'Atención manual inicial; por sí sola no integra nuestra IA personalizada con Yaku.'),
], [1.5, 2.05, 3.35])
sources([
    ('Meta acceso y precios', 'https://whatsappbusiness.com/resources/faq/'),
    ('Twilio tarifas', 'https://www.twilio.com/en-us/whatsapp/pricing'),
    ('360dialog planes', 'https://docs.360dialog.com/docs/get-started/pricing'),
])
h('Continuidad del número actual')
p('360dialog documenta coexistencia para cuentas Business activas y elegibles. La antigüedad y calidad influyen; no se puede prometer para cualquier número. Requiere mantener instalada la app y abrirla al menos cada 13 días. Algunos dispositivos vinculados no sincronizan con la API, por lo que probaremos el flujo que usa OLBOL.')
p('En Meta directa debemos revisar el flujo habilitado para nuestra aplicación. En Twilio, el alta estándar requiere un número no registrado o una migración planificada; no daremos por confirmada la coexistencia. No se debe liberar el número comercial antes de comprobar elegibilidad y respaldar lo necesario.')
sources([
    ('360dialog coexistencia', 'https://docs.360dialog.com/docs/resources/phone-numbers/coexistence'),
    ('Twilio alta estándar', 'https://www.twilio.com/docs/whatsapp/self-sign-up'),
])
h('Elección del CRM')
p('Yaku evita proponer una licencia de CRM ajena a nuestra plataforma. Su precio, usuarios incluidos y capacidades se deben definir. El CRM del DMS requiere desarrollo, persistencia, permisos y bandeja de atención si los vendedores responderán desde allí; un tablero de oportunidades por sí solo no cumple esa función.')

page('Requisitos para implementar cada ruta')
h('OLBOL debe aportar')
p('Acceso administrador al portafolio empresarial de Meta y una cuenta de WhatsApp Business o WABA; titularidad del número y recepción del código por SMS o llamada; nombre comercial, web y política de privacidad; medio de pago y documentación de empresa cuando el proceso de verificación la solicite. Confirmaremos los activos y su estado antes del alta.')
p('Para el piloto necesitamos catálogo aprobado, precios y disponibilidad actualizados, horarios, responsables, condiciones de venta y permiso de contacto para seguimientos. Definiremos qué información puede responder la IA y qué decisiones reserva OLBOL al vendedor.')
h('Meta Cloud API directa')
p('Crear o configurar la app de Meta, registrar el número en la WABA, generar un token de servidor y conceder whatsapp_business_messaging y whatsapp_business_management. business_management se utiliza cuando se consultan los activos empresariales. Suscribir la app a webhooks y desplegar un endpoint HTTPS. El token temporal de prueba no sirve para operar el servicio permanentemente.')
sources([('Colección oficial de Meta y requisitos Cloud API', 'https://www.postman.com/meta/whatsapp-business-platform/documentation/wlk6lh4/whatsapp-cloud-api')])
h('Twilio')
p('Cuenta Twilio de pago, administrador de los activos de Meta, número compatible y alta de un remitente mediante Self Sign-up. Su flujo puede exigir verificación empresarial para producción. Conectar las credenciales de servidor, webhook entrante y estados de entrega. El alquiler de un número Twilio y otros productos tienen costos propios.')
p('Para conectar clientes con Twilio como proveedor de software Yaku, debemos seguir el programa Tech Provider: app aprobada por Meta, vinculación Partner Solution e integración de Embedded Signup y Senders API. Twilio indica 3 a 4 semanas para los dos primeros pasos; el desarrollo añade tiempo. Self Sign-up es el flujo de cliente directo, no un sustituto del alta de clientes de Yaku.')
sources([
    ('Twilio Self Sign-up', 'https://www.twilio.com/docs/whatsapp/self-sign-up'),
    ('Twilio Tech Provider', 'https://www.twilio.com/docs/whatsapp/isv/tech-provider-program'),
])
h('360dialog y validación de Yaku')
p('Cuenta Hub, plan y pago, Embedded Signup con administrador de Meta, WABA y verificación del número. Elegir alta estándar o coexistencia según elegibilidad; configurar API key y webhook. El plan incluye el canal, no nuestra bandeja CRM ni el asesor personalizado.')
p('Antes de cotizar, comprobar en Yaku API y webhooks, lectura y escritura de mensajes, contactos y oportunidades, asignación, pausa del bot y separación de clientes. Para ofrecer onboarding de varias empresas desde Yaku también validaremos los requisitos de Meta para proveedores tecnológicos. Son puntos por verificar, no capacidades confirmadas.')
sources([('360dialog Embedded Signup', 'https://docs.360dialog.com/docs/hub/embedded-signup')])

page('Cómo atenderá la IA y cuándo pasa al vendedor')
p('El mensaje llega por el webhook oficial. Nuestro backend verifica el evento, evita duplicados y registra el historial. Consulta el catálogo vigente, genera una respuesta acotada a OLBOL y la envía por la API elegida. Yaku o el DMS reciben el contacto, el interés, el origen y la próxima acción. Este es el diseño propuesto para producción.')
shot('recomendacion-escena.png', 'Asesor web de demostración. El piloto llevará este tipo de orientación al canal oficial de WhatsApp con el catálogo aprobado de OLBOL.', 5.7)
h('Control humano y alcance inicial')
p('La IA orienta sobre modelos y requisitos, solicita datos mínimos y propone una visita. No inventa stock ni precios y no aprueba crédito ni descuentos. Ante una solicitud humana, una negociación o información incierta, asigna un vendedor y pausa el bot para evitar respuestas simultáneas. Incluimos reintentos, registro de errores, permisos y alertas como requisitos de ingeniería.')
h('Reducir el riesgo de restricciones')
p('La política permite automatización con acceso claro a atención humana. Respetaremos consentimiento para contactos posteriores, bajas y plantillas aprobadas al iniciar conversaciones o responder fuera de las 24 horas desde el último mensaje del cliente. Registraremos permisos y revisaremos calidad y restricciones. La API oficial no garantiza cero bloqueos. El asistente se limita a ventas y consultas de OLBOL; validaremos los términos aplicables antes de activarlo.')
sources([('Política oficial de WhatsApp', 'https://whatsappbusiness.com/policy/')])

page('Costos de canal y consumo de IA')
p('Escenarios calculados para un número: cada consulta tiene 5 mensajes del cliente y 5 respuestas de texto de la IA. Suponemos 2.000 tokens de entrada y 200 de salida por respuesta, sin caché. GPT 4.1 mini se usa como referencia de costo: USD 0,40 por millón de tokens de entrada y USD 1,60 de salida. El modelo final se elige con pruebas de calidad.')
table(['Consultas al mes', 'Mensajes totales', 'IA texto estimada', 'Meta acceso', 'Twilio recargo', '360dialog cuota'], [
    ('500', '5.000', 'USD 2,80', 'USD 0', 'USD 25', 'USD 59'),
    ('1.500', '15.000', 'USD 8,40', 'USD 0', 'USD 75', 'USD 59'),
    ('5.000', '50.000', 'USD 28', 'USD 0', 'USD 250', 'USD 59'),
], [1.1, 1.1, 1.3, 1.0, 1.2, 1.2], centered=(0, 1, 2, 3, 4, 5))
p('Estas columnas no son totales. Falta sumar consumo de Meta, infraestructura, CRM y honorarios. No incluyen audios, imágenes, herramientas, reintentos, impuestos, números alquilados ni complementos. Cálculo propio: cada respuesta consume USD 0,00112 en este supuesto.')
sources([('OpenAI precios del modelo de referencia', 'https://developers.openai.com/api/docs/models/gpt-4.1-mini')])
h('Meta se presupuesta por separado')
p('Desde el 1 de octubre de 2026, Meta publica cargos para mensajes de atención, con los primeros 1.000 por número al mes sin cargo. Las plantillas tienen su categoría y tarifa por mercado. La ventana de atención determina qué se puede enviar; no implica gratuidad ilimitada. Confirmaremos la tarifa unitaria vigente para Bolivia, el destino de los contactos y las exenciones aplicables antes de emitir oferta.')
sources([('Meta actualización y precios', 'https://whatsappbusiness.com/resources/faq/')])
h('Presupuesto para discutir en la demo')
table(['Concepto', 'Referencia de planificación'], [
    ('Implementación y mantenimiento', 'USD 2.400 iniciales y USD 250 al mes sugeridos. Revalidar tras auditar Yaku y definir alcance.'),
    ('Infraestructura e IA', 'Reserva inicial de USD 25 al mes para cada uno. No es una tarifa ni un tope garantizado.'),
    ('Meta anuncios y CRM', 'Meta: reserva orientativa USD 30 a 90, pendiente del tarifario. Anuncios: ejemplo USD 450. Yaku o desarrollo del CRM DMS: por cotizar.'),
], [2.0, 4.9])
p('Con 10.000 mensajes al mes, estas reservas y honorarios, el subtotal es USD 780 a 840 con Meta directa, USD 830 a 890 con Twilio o USD 839 a 899 con 360dialog. Sumar CRM y exclusiones. A 11.800 mensajes, el recargo Twilio iguala USD 59; este cruce compara solo el proveedor, no soporte ni costo de ingeniería.')

page('Qué presentar y qué acordar con OLBOL')
h('Speech de presentación de la solución')
p('“Queremos que cada consulta reciba una respuesta útil y tenga un responsable que la acompañe hasta la visita o la compra. Proponemos Yaku para organizar ese seguimiento y una IA que responda en WhatsApp con la información aprobada de su catálogo. Cuando el comprador necesita negociar o hablar con alguien, su vendedor toma el control.”')
p('“Para conectarnos usamos la plataforma oficial de WhatsApp. Nuestra primera opción es Meta directa; también podemos elegir un proveedor si facilita conservar el uso del número o la operación. Vamos a validar su cuenta y explicar los cargos antes de activar el piloto. Nadie puede garantizar que Meta nunca restrinja una cuenta; nuestro compromiso es trabajar con permisos, buenas prácticas y supervisión.”')
p('“Si prefiere tener el CRM dentro de su panel administrativo, podemos cotizar esa opción. Hoy acordemos el flujo que necesita y un piloto con un número, un responsable y una campaña. Así verificamos que las consultas llegan, se atienden y avanzan, antes de ampliar la inversión.”')
h('Recorrido de la demo en 10 minutos')
bullet('Catálogo y ficha: mostrar exploración móvil y seleccionar Terra S5. Recordar que precios y stock de la demo son ficticios.')
bullet('Asesor: escribir “Somos 5, viajamos los fines de semana y tengo un presupuesto de USD 30 mil”. Pedir cotización y preparar continuidad.')
bullet('Consola: abrir la oportunidad, asignar responsable y próxima acción. Explicar que representa el flujo comercial, no el Yaku real.')
bullet('Propuesta: elegir Yaku o DMS, comparar canales y ajustar mensajes y anuncios. Presentar el subtotal con el CRM pendiente, sin dar un precio cerrado.')
h('Condiciones de aceptación del piloto')
p('Proponemos probar mensajes entrantes y salientes, estados de entrega, generación de oportunidades, asignación humana y pausa de IA. Una baja debe detener el seguimiento; un reintento no debe duplicar mensajes. Evaluaremos respuestas contra el catálogo, casos sin respuesta y atención fuera de horario. Los objetivos comerciales se fijan con el volumen y la situación inicial de OLBOL.')
h('Cierre propuesto')
p('“El siguiente paso es definir el CRM elegido y entregarnos catálogo, acceso a sus activos y un responsable comercial. Con esa revisión le presentamos el alcance y precio final del piloto. Si el flujo cumple lo acordado, activamos la campaña y medimos resultados con usted.”')
p('El calendario de 4 a 6 semanas es orientativo: depende de la auditoría de Yaku, alcance del DMS y aprobaciones del proveedor y Meta. La demo actual no envía WhatsApp ni tiene campañas activas. Licencia, número de vendedores y capacidad técnica de Yaku deben quedar resueltos en la oferta final.')

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
